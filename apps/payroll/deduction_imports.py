"""Bulk ADD of employee deductions from an Excel template.

Mirrors apps.employees.imports.build_create_plan / apply_create_plan: always
creates new EmployeeDeduction rows (never updates an existing one - an
employee can legitimately carry more than one deduction of the same code,
e.g. a second loan after the first is paid off). Matches by Staff ID and
deduction code, validates each row the same way the on-screen form does
(EmployeeDeduction.full_clean()), and commits nothing until confirmed.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.exceptions import ValidationError
from django.db import transaction

from apps.employees.models import Employee

from .models import DeductionCode, DeductionScope, EmployeeDeduction, Recurrence

REQUIRED = ["staff_id", "deduction_code", "amount"]
OPTIONAL = [
    "run_scope", "recurrence", "opening_balance", "number_of_runs",
    "effective_from", "effective_to", "note",
]
TEMPLATE_HEADER = REQUIRED + OPTIONAL


@dataclass
class DeductionRow:
    line_no: int
    staff_id: str
    code: str
    error: str = ""
    instance: EmployeeDeduction | None = None

    @property
    def ok(self):
        return not self.error


@dataclass
class DeductionImportPlan:
    rows: list = field(default_factory=list)
    header_error: str = ""

    @property
    def valid_rows(self):
        return [r for r in self.rows if r.ok]

    @property
    def error_rows(self):
        return [r for r in self.rows if r.error]


def _parse_date(raw):
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    raw = str(raw).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{raw}' is not a recognised date (use YYYY-MM-DD)")


def _parse_decimal(raw, field_name):
    if raw in (None, ""):
        return None
    if isinstance(raw, (int, float, Decimal)):
        val = Decimal(str(raw))
    else:
        try:
            val = Decimal(str(raw).strip().replace(",", ""))
        except InvalidOperation:
            raise ValueError(f"{field_name}: '{raw}' is not a number")
    if val < 0:
        raise ValueError(f"{field_name}: must not be negative")
    return val.quantize(Decimal("0.01"))


def _parse_int(raw, field_name):
    if raw in (None, ""):
        return None
    try:
        val = int(float(raw))
    except (TypeError, ValueError):
        raise ValueError(f"{field_name}: '{raw}' is not a whole number")
    if val < 1:
        raise ValueError(f"{field_name}: must be at least 1")
    return val


def _rows_from_xlsx(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        return [], []
    cols = [(c or "").strip().lower() for c in header]
    rows = []
    for raw in it:
        if raw is None or all(v in (None, "") for v in raw):
            continue
        rows.append(dict(zip(cols, raw)))
    return cols, rows


def build_plan(file_obj) -> DeductionImportPlan:
    cols, raw_rows = _rows_from_xlsx(file_obj)
    plan = DeductionImportPlan()

    missing = [c for c in REQUIRED if c not in cols]
    if missing:
        plan.header_error = "Missing required column(s): " + ", ".join(missing)
        return plan

    employees = {e.staff_id: e for e in Employee.objects.all()}
    codes = {c.code.upper(): c for c in DeductionCode.objects.filter(is_active=True)}

    for i, row in enumerate(raw_rows, start=2):
        sid = str(row.get("staff_id") or "").strip()
        code_raw = str(row.get("deduction_code") or "").strip()
        r = DeductionRow(line_no=i, staff_id=sid, code=code_raw)

        emp = employees.get(sid)
        if not sid:
            r.error = "missing staff_id"
            plan.rows.append(r)
            continue
        if emp is None:
            r.error = "no employee with this Staff ID"
            plan.rows.append(r)
            continue

        code = codes.get(code_raw.upper())
        if not code_raw:
            r.error = "missing deduction_code"
            plan.rows.append(r)
            continue
        if code is None:
            r.error = f"unknown deduction_code '{code_raw}'"
            plan.rows.append(r)
            continue

        try:
            amount = _parse_decimal(row.get("amount"), "amount")
            if amount is None:
                raise ValueError("amount: required")
            run_scope = str(row.get("run_scope") or "").strip().upper() or DeductionScope.BOTH
            if run_scope not in DeductionScope.values:
                raise ValueError(f"run_scope: must be one of {', '.join(DeductionScope.values)}")
            recurrence = str(row.get("recurrence") or "").strip().upper() or Recurrence.RECURRING
            if recurrence not in Recurrence.values:
                raise ValueError(f"recurrence: must be one of {', '.join(Recurrence.values)}")

            ed = EmployeeDeduction(
                employee=emp,
                code=code,
                amount=amount,
                run_scope=run_scope,
                recurrence=recurrence,
                opening_balance=_parse_decimal(row.get("opening_balance"), "opening_balance"),
                number_of_runs=_parse_int(row.get("number_of_runs"), "number_of_runs"),
                effective_from=_parse_date(row.get("effective_from")),
                effective_to=_parse_date(row.get("effective_to")),
                note=str(row.get("note") or "").strip(),
            )
            ed.full_clean(exclude=["balance", "runs_applied"])
            r.instance = ed
        except ValidationError as exc:
            r.error = "; ".join(f"{k}: {' '.join(v)}" for k, v in exc.message_dict.items())
        except ValueError as exc:
            r.error = str(exc)
        plan.rows.append(r)

    return plan


@transaction.atomic
def apply_plan(plan: DeductionImportPlan) -> dict:
    created = 0
    for r in plan.valid_rows:
        r.instance.save()
        created += 1
    return {"created": created, "skipped": len(plan.error_rows)}


def template_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deductions"

    bold = openpyxl.styles.Font(bold=True)
    ws.append(TEMPLATE_HEADER)
    for cell in ws[1]:
        cell.font = bold

    sample_emp = Employee.objects.exclude(staff_id="").first()
    sample_code = DeductionCode.objects.filter(is_active=True, is_loan=False).first()
    if sample_emp and sample_code:
        ws.append([
            sample_emp.staff_id, sample_code.code, "50000.00",
            "BOTH", "RECURRING", "", "", "", "", "e.g. monthly contribution",
        ])
    for col_idx, header in enumerate(TEMPLATE_HEADER, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(14, len(header) + 2)

    ref = wb.create_sheet("Reference")
    ref.append(["Deduction codes", "Name", "Is loan (needs opening_balance)", "Applies to Salary run", "Applies to Expense Allowance run"])
    for cell in ref[1]:
        cell.font = bold
    for c in DeductionCode.objects.filter(is_active=True).order_by("sequence", "code"):
        ref.append([c.code, c.name, "Yes" if c.is_loan else "No", "Yes" if c.applies_to_salary else "No", "Yes" if c.applies_to_expense else "No"])
    ref.append([])
    ref.append(["run_scope choices:", ", ".join(DeductionScope.values)])
    ref.append(["recurrence choices:", ", ".join(Recurrence.values)])
    ref.append(["Dates:", "use YYYY-MM-DD"])
    for col_idx in range(1, 6):
        ref.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
